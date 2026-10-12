import os, io, csv, json, sqlite3, hashlib, secrets, time, urllib.parse
from datetime import datetime, timedelta
from functools import wraps
from flask import Flask, request, session, jsonify, render_template, redirect, send_file, g
from flask_cors import CORS
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas
import qrcode
from PIL import Image, ImageDraw, ImageFont

BASE = os.path.dirname(os.path.abspath(__file__))
DB = os.environ.get("DB_PATH", os.path.join(BASE, "sunu.db"))
BASE_URL = os.environ.get("BASE_URL", "https://factuor.onrender.com")
SALT = os.environ.get("PW_SALT", "sunu-jang-pay")
ADMIN_KEY = os.environ.get("ADMIN_KEY", "change-moi")
WAVE_OM = "785278668"
LOGO = os.path.join(BASE, "static", "logo.png")
ICON = os.path.join(BASE, "static", "logo_icon.png")  # remplacé plus bas par la version intégrée

import base64
from assets import LOGO_B64, ICON_B64
ICON = lambda: io.BytesIO(base64.b64decode(ICON_B64))
from flask import Response
app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "dev-secret-change-me")
CORS(app)

SCHEMA = """
CREATE TABLE IF NOT EXISTS ecoles(id INTEGER PRIMARY KEY, nom TEXT, email TEXT UNIQUE, password_hash TEXT, tel TEXT, date_inscription TEXT, abonne_jusqu TEXT);
CREATE TABLE IF NOT EXISTS eleves(id INTEGER PRIMARY KEY, ecole_id INTEGER, matricule TEXT UNIQUE, nom TEXT, classe TEXT, parent_nom TEXT, parent_tel TEXT, parent_tel2 TEXT, montant_base INTEGER DEFAULT 0, statut TEXT DEFAULT 'EN ATTENTE', libelles TEXT DEFAULT '[]');
CREATE INDEX IF NOT EXISTS ix_el ON eleves(ecole_id, matricule);
CREATE INDEX IF NOT EXISTS ix_el_cl ON eleves(ecole_id, classe, statut);
CREATE TABLE IF NOT EXISTS libelles(id INTEGER PRIMARY KEY, ecole_id INTEGER, nom TEXT, montant INTEGER);
CREATE TABLE IF NOT EXISTS paiements(id INTEGER PRIMARY KEY, ecole_id INTEGER, eleve_id INTEGER, recu TEXT UNIQUE, libelle TEXT, mois TEXT, mode TEXT, montant INTEGER, date TEXT, par TEXT);
CREATE INDEX IF NOT EXISTS ix_pay ON paiements(ecole_id, date);
CREATE TABLE IF NOT EXISTS historique(id INTEGER PRIMARY KEY, ecole_id INTEGER, action TEXT, details TEXT, date TEXT, ip TEXT);
CREATE INDEX IF NOT EXISTS ix_hist ON historique(ecole_id, id);
"""

def db():
    if "db" not in g:
        g.db = sqlite3.connect(DB, timeout=30)
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA journal_mode=WAL"); g.db.execute("PRAGMA synchronous=NORMAL")
    return g.db

@app.teardown_appcontext
def close(_):
    d = g.pop("db", None)
    if d: d.close()

with sqlite3.connect(DB) as _c:
    _c.executescript(SCHEMA)

now = lambda: datetime.now().isoformat(timespec="seconds")
h = lambda p: hashlib.sha256((SALT + p).encode()).hexdigest()
ip = lambda: request.headers.get("X-Forwarded-For", request.remote_addr or "").split(",")[0]

def log(eid, action, details=""):
    db().execute("INSERT INTO historique(ecole_id,action,details,date,ip) VALUES(?,?,?,?,?)", (eid, action, details, now(), ip()))

_hits = {}
def limited(key, n=10, win=60):
    t = time.time(); l = [x for x in _hits.get(key, []) if t - x < win]
    l.append(t); _hits[key] = l
    return len(l) > n

def prix(nb): return 15000 if nb <= 200 else 30000 if nb <= 500 else 45000

def etat(eid):
    d = db(); e = d.execute("SELECT * FROM ecoles WHERE id=?", (eid,)).fetchone()
    nb = d.execute("SELECT COUNT(*) FROM eleves WHERE ecole_id=?", (eid,)).fetchone()[0]
    fin_essai = datetime.fromisoformat(e["date_inscription"]) + timedelta(days=7)
    paye = bool(e["abonne_jusqu"]) and datetime.fromisoformat(e["abonne_jusqu"]) > datetime.now()
    essai = datetime.now() < fin_essai
    return dict(nb=nb, prix=prix(nb), actif=paye or essai, essai=essai and not paye,
                jours_essai=max(0, (fin_essai - datetime.now()).days + 1) if essai else 0, ecole=e)

def need(page=False, sub=True):
    def deco(f):
        @wraps(f)
        def w(*a, **k):
            eid = session.get("ecole_id")
            if not eid:
                return redirect("/") if page else (jsonify(error="Non connecté"), 401)
            if "ecole_id" in k and int(k["ecole_id"]) != eid:
                return jsonify(error="Accès refusé"), 403
            if sub and not etat(eid)["actif"]:
                return redirect("/abonnement") if page else (jsonify(error="Abonnement expiré"), 402)
            return f(*a, **k)
        return w
    return deco

# ---------- Logos (intégrés, toujours disponibles) ----------
@app.get("/static/logo.png")
def logo_png():
    return Response(base64.b64decode(LOGO_B64), mimetype="image/png", headers={"Cache-Control": "public, max-age=86400"})

@app.get("/static/logo_icon.png")
def logo_icon():
    return Response(base64.b64decode(ICON_B64), mimetype="image/png", headers={"Cache-Control": "public, max-age=86400"})

# ---------- Pages ----------
@app.route("/")
def login_page():
    return redirect("/dashboard") if session.get("ecole_id") else render_template("login.html")

@app.route("/dashboard")
@need(page=True)
def dashboard():
    s = etat(session["ecole_id"])
    return render_template("dashboard.html", ecole=s["ecole"]["nom"], eid=session["ecole_id"], essai=s["essai"], jours=s["jours_essai"])

@app.route("/abonnement")
@need(page=True, sub=False)
def abonnement():
    s = etat(session["ecole_id"])
    return render_template("abonnement.html", ecole=s["ecole"]["nom"], nb=s["nb"], prix=s["prix"], actif=s["actif"], numero=WAVE_OM,
                           wa="https://wa.me/221" + WAVE_OM + "?text=" + urllib.parse.quote(f"Bonjour, j'ai payé l'abonnement SUNU JÀNG PAY pour {s['ecole']['nom']} ({s['ecole']['email']}) : {s['prix']} FCFA."))

@app.route("/logout")
def logout():
    session.clear(); return redirect("/")

# ---------- Auth ----------
@app.post("/api/register")
def register():
    if limited("reg" + ip(), 5, 300): return jsonify(error="Trop de tentatives"), 429
    j = request.json or {}
    nom, email, pw, tel = (j.get(k, "").strip() for k in ("nom", "email", "password", "tel"))
    if not (nom and email and len(pw) >= 6): return jsonify(error="Nom, email et mot de passe (6+ car.) requis"), 400
    d = db()
    try:
        c = d.execute("INSERT INTO ecoles(nom,email,password_hash,tel,date_inscription) VALUES(?,?,?,?,?)", (nom, email.lower(), h(pw), tel, now()))
    except sqlite3.IntegrityError:
        return jsonify(error="Cet email existe déjà"), 409
    log(c.lastrowid, "INSCRIPTION", nom); d.commit()
    session["ecole_id"] = c.lastrowid
    return jsonify(ok=True)

@app.post("/api/login")
def login():
    if limited("log" + ip()): return jsonify(error="Trop de tentatives, réessayez dans 1 minute"), 429
    j = request.json or {}
    e = db().execute("SELECT id FROM ecoles WHERE email=? AND password_hash=?", ((j.get("email") or "").lower().strip(), h(j.get("password") or ""))).fetchone()
    if not e: return jsonify(error="Email ou mot de passe incorrect"), 401
    session["ecole_id"] = e["id"]; log(e["id"], "CONNEXION"); db().commit()
    return jsonify(ok=True)

@app.get("/api/admin/activer")  # /api/admin/activer?key=...&ecole=1&jours=30
def activer():
    if request.args.get("key") != ADMIN_KEY: return jsonify(error="Clé invalide"), 403
    eid = int(request.args["ecole"]); fin = (datetime.now() + timedelta(days=int(request.args.get("jours", 30)))).isoformat(timespec="seconds")
    db().execute("UPDATE ecoles SET abonne_jusqu=? WHERE id=?", (fin, eid)); log(eid, "ABONNEMENT ACTIVÉ", fin); db().commit()
    return jsonify(ok=True, jusqua=fin)

# ---------- Élèves ----------
def new_matricules(n):
    ex = {r[0] for r in db().execute("SELECT matricule FROM eleves")}
    out = []
    while len(out) < n:
        m = f"SJ-{datetime.now().year}-{secrets.token_hex(3).upper()}"
        if m not in ex: ex.add(m); out.append(m)
    return out

def total_du(e, lib):
    return (e["montant_base"] or 0) + sum(lib.get(i, 0) for i in json.loads(e["libelles"] or "[]"))

def libmap(eid): return {r["id"]: r["montant"] for r in db().execute("SELECT id,montant FROM libelles WHERE ecole_id=?", (eid,))}

@app.get("/api/eleves/<int:ecole_id>")
@need()
def eleves(ecole_id):
    q, cl, st = request.args.get("q", "").strip(), request.args.get("classe", ""), request.args.get("statut", "")
    page = max(1, int(request.args.get("page", 1)))
    w, p = ["ecole_id=?"], [ecole_id]
    if q: w.append("(nom LIKE ? OR matricule LIKE ? OR parent_nom LIKE ? OR parent_tel LIKE ?)"); p += [f"%{q}%"] * 4
    if cl: w.append("classe=?"); p.append(cl)
    if st: w.append("statut=?"); p.append(st)
    W = " AND ".join(w); d = db()
    total = d.execute(f"SELECT COUNT(*) FROM eleves WHERE {W}", p).fetchone()[0]
    rows = d.execute(f"SELECT * FROM eleves WHERE {W} ORDER BY nom LIMIT 100 OFFSET ?", p + [(page - 1) * 100]).fetchall()
    lib = libmap(ecole_id)
    out = [dict(r, libelles=json.loads(r["libelles"]), total=total_du(r, lib)) for r in rows]
    classes = [r[0] for r in d.execute("SELECT DISTINCT classe FROM eleves WHERE ecole_id=? ORDER BY classe", (ecole_id,))]
    return jsonify(eleves=out, total=total, pages=(total + 99) // 100, page=page, classes=classes)

@app.post("/api/eleves/add")
@need()
def add_eleve():
    j = request.json or {}; eid = session["ecole_id"]
    if not j.get("nom"): return jsonify(error="Nom requis"), 400
    m = new_matricules(1)[0]
    db().execute("INSERT INTO eleves(ecole_id,matricule,nom,classe,parent_nom,parent_tel,parent_tel2,montant_base,libelles) VALUES(?,?,?,?,?,?,?,?,?)",
                 (eid, m, j["nom"].strip(), j.get("classe", "").strip(), j.get("parent_nom", ""), j.get("parent_tel", ""), j.get("parent_tel2", ""), int(j.get("montant_base") or 0), json.dumps(j.get("libelles", []))))
    log(eid, "ÉLÈVE AJOUTÉ", f"{j['nom']} {m}"); db().commit()
    return jsonify(ok=True, matricule=m)

@app.post("/api/eleves/import")  # CSV ou XLSX : nom,classe,parent_nom,parent_tel,parent_tel2,montant_base
@need()
def import_eleves():
    f = request.files.get("file"); eid = session["ecole_id"]
    if not f: return jsonify(error="Fichier manquant"), 400
    if f.filename.lower().endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook
        except ImportError:
            return jsonify(error="Installez openpyxl (pip install openpyxl) ou envoyez un CSV"), 400
        rows = [[("" if c is None else str(c)) for c in r] for r in load_workbook(f, read_only=True).active.iter_rows(values_only=True)]
    else:
        txt = f.read().decode("utf-8-sig", "ignore"); rows = list(csv.reader(io.StringIO(txt), delimiter=";" if txt.count(";") > txt.count(",") else ","))
    if not rows: return jsonify(error="Fichier vide"), 400
    head = [c.strip().lower() for c in rows[0]]
    if "nom" in head: rows = rows[1:]
    else: head = ["nom", "classe", "parent_nom", "parent_tel", "parent_tel2", "montant_base"]
    data = [dict(zip(head, r)) for r in rows if r and any(r)]
    mats = new_matricules(len(data)); batch = []
    for m, r in zip(mats, data):
        try: mb = int(float(str(r.get("montant_base") or 0).replace(" ", "")))
        except ValueError: mb = 0
        batch.append((eid, m, r.get("nom", "").strip(), r.get("classe", "").strip(), r.get("parent_nom", ""), r.get("parent_tel", ""), r.get("parent_tel2", ""), mb))
    db().executemany("INSERT INTO eleves(ecole_id,matricule,nom,classe,parent_nom,parent_tel,parent_tel2,montant_base) VALUES(?,?,?,?,?,?,?,?)", batch)
    log(eid, "IMPORT", f"{len(batch)} élèves"); db().commit()
    return jsonify(ok=True, n=len(batch))

@app.post("/api/eleves/libelles")
@need()
def eleve_libelles():
    j = request.json or {}
    db().execute("UPDATE eleves SET libelles=? WHERE id=? AND ecole_id=?", (json.dumps(j.get("libelles", [])), j.get("eleve_id"), session["ecole_id"]))
    db().commit(); return jsonify(ok=True)

@app.route("/api/libelles", methods=["GET", "POST", "DELETE"])
@need()
def libelles():
    eid = session["ecole_id"]; d = db()
    if request.method == "POST":
        j = request.json or {}
        if not j.get("nom"): return jsonify(error="Nom requis"), 400
        d.execute("INSERT INTO libelles(ecole_id,nom,montant) VALUES(?,?,?)", (eid, j["nom"].strip(), int(j.get("montant") or 0)))
        log(eid, "LIBELLÉ AJOUTÉ", f"{j['nom']} {j.get('montant')}"); d.commit()
    if request.method == "DELETE":
        d.execute("DELETE FROM libelles WHERE id=? AND ecole_id=?", (request.args.get("id"), eid)); d.commit()
    return jsonify(libelles=[dict(r) for r in d.execute("SELECT * FROM libelles WHERE ecole_id=? ORDER BY nom", (eid,))])

# ---------- Encaissement ----------
@app.post("/api/payer")
@need()
def payer():
    j = request.json or {}; eid = session["ecole_id"]; d = db()
    e = d.execute("SELECT * FROM eleves WHERE id=? AND ecole_id=?", (j.get("eleve_id"), eid)).fetchone()
    if not e: return jsonify(error="Élève introuvable"), 404
    mode = j.get("mode", "Cash")
    if mode not in ("Wave", "Orange Money", "Cash"): return jsonify(error="Mode invalide"), 400
    montant = int(j.get("montant") or 0)
    if montant <= 0: return jsonify(error="Montant invalide"), 400
    recu = "RECU-" + secrets.token_hex(4).upper()
    d.execute("INSERT INTO paiements(ecole_id,eleve_id,recu,libelle,mois,mode,montant,date,par) VALUES(?,?,?,?,?,?,?,?,?)",
              (eid, e["id"], recu, j.get("libelle", "Scolarité"), j.get("mois", ""), mode, montant, now(), session.get("ecole_id")))
    d.execute("UPDATE eleves SET statut='PAYÉ' WHERE id=?", (e["id"],))
    log(eid, "ENCAISSEMENT", f"{e['nom']} {e['matricule']} {montant} FCFA {mode} {j.get('libelle','')} {j.get('mois','')} [{recu}]"); d.commit()
    return jsonify(ok=True, recu=recu)

@app.post("/api/statut")
@need()
def statut():
    j = request.json or {}
    if j.get("statut") not in ("PAYÉ", "RETARD", "EN ATTENTE"): return jsonify(error="Statut invalide"), 400
    db().execute("UPDATE eleves SET statut=? WHERE id=? AND ecole_id=?", (j["statut"], j.get("eleve_id"), session["ecole_id"]))
    log(session["ecole_id"], "STATUT", f"élève {j.get('eleve_id')} → {j['statut']}"); db().commit(); return jsonify(ok=True)

@app.get("/api/stats/<int:ecole_id>")
@need()
def stats(ecole_id):
    d = db(); mois = datetime.now().strftime("%Y-%m")
    enc = d.execute("SELECT COALESCE(SUM(montant),0) FROM paiements WHERE ecole_id=? AND date LIKE ?", (ecole_id, mois + "%")).fetchone()[0]
    r = d.execute("SELECT COUNT(*) t, SUM(statut='PAYÉ') p, SUM(statut='RETARD') r, COALESCE(SUM(CASE WHEN statut='RETARD' THEN montant_base END),0) mr FROM eleves WHERE ecole_id=?", (ecole_id,)).fetchone()
    t, p = r["t"], r["p"] or 0
    ev = {x["m"]: x["s"] for x in d.execute("SELECT substr(date,1,7) m, SUM(montant) s FROM paiements WHERE ecole_id=? GROUP BY m", (ecole_id,))}
    ms = []; y, m = datetime.now().year, datetime.now().month
    for _ in range(6): ms.append(f"{y}-{m:02d}"); m -= 1; y, m = (y - 1, 12) if m == 0 else (y, m)
    ms.reverse()
    return jsonify(encaisse=enc, retard=r["r"] or 0, montant_retard=r["mr"], payes=p, total=t, taux=round(100 * p / t) if t else 0,
                   evolution=dict(labels=ms, values=[ev.get(x, 0) for x in ms]))

@app.get("/api/caisse")
@need()
def caisse():
    rows = db().execute("SELECT p.*, e.nom, e.classe FROM paiements p JOIN eleves e ON e.id=p.eleve_id WHERE p.ecole_id=? AND p.date LIKE ? ORDER BY p.id DESC", (session["ecole_id"], datetime.now().strftime("%Y-%m-%d") + "%")).fetchall()
    tot = {"Cash": 0, "Wave": 0, "Orange Money": 0}
    for r in rows: tot[r["mode"]] += r["montant"]
    return jsonify(totaux=tot, total=sum(tot.values()), paiements=[dict(r) for r in rows])

@app.get("/api/historique")
@need()
def historique():
    return jsonify([dict(r) for r in db().execute("SELECT action,details,date,ip FROM historique WHERE ecole_id=? ORDER BY id DESC LIMIT 200", (session["ecole_id"],))])

@app.get("/api/parents")
@need()
def parents():
    eid = session["ecole_id"]; rows = db().execute("SELECT matricule,nom,classe,parent_nom,parent_tel,parent_tel2 FROM eleves WHERE ecole_id=? ORDER BY classe,nom LIMIT 100 OFFSET ?", (eid, (max(1, int(request.args.get("page", 1))) - 1) * 100)).fetchall()
    if request.args.get("export"):
        s = io.StringIO(); w = csv.writer(s, delimiter=";"); w.writerow(["Matricule", "Élève", "Classe", "Parent", "Tel", "Tel 2"])
        for r in db().execute("SELECT matricule,nom,classe,parent_nom,parent_tel,parent_tel2 FROM eleves WHERE ecole_id=? ORDER BY classe,nom", (eid,)): w.writerow(list(r))
        return send_file(io.BytesIO(("\ufeff" + s.getvalue()).encode()), mimetype="text/csv", as_attachment=True, download_name="parents.csv")
    return jsonify([dict(r) for r in rows])

# ---------- Reçus ----------
def recu_data(code):
    r = db().execute("SELECT p.*, e.nom, e.classe, e.matricule, e.parent_nom, s.nom ecole FROM paiements p JOIN eleves e ON e.id=p.eleve_id JOIN ecoles s ON s.id=p.ecole_id WHERE p.recu=?", (code,)).fetchone()
    return r

fcfa = lambda n: f"{n:,}".replace(",", " ") + " FCFA"

def recu_pdf(r):
    buf = io.BytesIO(); c = canvas.Canvas(buf, pagesize=A4); W, H = A4
    c.setFillColorRGB(.059, .09, .165); c.rect(0, H - 170, W, 170, fill=1, stroke=0)
    c.drawImage(ImageReader(ICON()), 40, H - 150, 120, 120, mask="auto")
    c.setFillColorRGB(1, 1, 1); c.setFont("Helvetica-Bold", 22); c.drawString(180, H - 80, "SUNU JÀNG PAY")
    c.setFont("Helvetica", 11); c.drawString(180, H - 100, "Fini de courir après les mensualités."); c.drawString(180, H - 125, r["ecole"])
    c.setFillColorRGB(.1, .1, .1); c.setFont("Helvetica-Bold", 18); c.drawString(40, H - 215, "REÇU DE PAIEMENT")
    c.setFont("Helvetica", 10); c.setFillColorRGB(.4, .4, .4); c.drawRightString(W - 40, H - 215, r["recu"])
    y = H - 255
    for k, v in [("Élève", r["nom"]), ("Matricule", r["matricule"]), ("Classe", r["classe"]), ("Parent", r["parent_nom"]), ("Libellé", r["libelle"]), ("Mois", r["mois"]), ("Mode", r["mode"]), ("Date", r["date"].replace("T", " "))]:
        c.setFillColorRGB(.4, .4, .4); c.setFont("Helvetica", 11); c.drawString(40, y, k); c.setFillColorRGB(.05, .05, .05); c.setFont("Helvetica-Bold", 11); c.drawString(150, y, str(v or "-")); y -= 24
    c.setFillColorRGB(.388, .4, .945); c.roundRect(40, y - 50, W - 80, 50, 8, fill=1, stroke=0)
    c.setFillColorRGB(1, 1, 1); c.setFont("Courier-Bold", 20); c.drawString(55, y - 32, "MONTANT : " + fcfa(r["montant"]))
    qr = qrcode.make(f"{BASE_URL}/verify/{r['recu']}"); qb = io.BytesIO(); qr.save(qb, "PNG"); qb.seek(0)
    c.drawImage(ImageReader(qb), 40, 90, 110, 110)
    c.setFillColorRGB(.4, .4, .4); c.setFont("Helvetica", 8); c.drawString(40, 78, "Scannez pour vérifier ce reçu")
    c.saveState(); c.translate(320, 150); c.rotate(12); c.setStrokeColorRGB(.13, .77, .37); c.setFillColorRGB(.13, .77, .37); c.setLineWidth(3)
    c.circle(0, 0, 45); c.circle(0, 0, 40); c.setFont("Helvetica-Bold", 20); c.drawCentredString(0, -2, "PAYÉ"); c.setFont("Helvetica", 7); c.drawCentredString(0, -16, r["date"][:10]); c.restoreState()
    c.setStrokeColorRGB(.2, .2, .2); c.setLineWidth(.6); c.line(400, 120, W - 40, 120)
    c.setFillColorRGB(.2, .2, .2); c.setFont("Helvetica-Oblique", 16); c.drawCentredString(465, 128, "Le Directeur"); c.setFont("Helvetica", 8); c.drawCentredString(465, 108, "Signature – " + r["ecole"])
    c.showPage(); c.save(); buf.seek(0); return buf

def recu_png(r):
    im = Image.new("RGB", (800, 1000), "white"); d = ImageDraw.Draw(im)
    f = lambda s: ImageFont.load_default(size=s)
    d.rectangle((0, 0, 800, 200), fill="#0F172A"); ic = Image.open(ICON()).convert("RGBA").resize((140, 140)); im.paste(ic, (30, 30), ic)
    d.text((200, 60), "SUNU JÀNG PAY", font=f(40), fill="white"); d.text((200, 120), r["ecole"], font=f(24), fill="#94A3B8")
    d.text((40, 240), "REÇU " + r["recu"], font=f(30), fill="#0F172A"); y = 310
    for k, v in [("Élève", r["nom"]), ("Matricule", r["matricule"]), ("Classe", r["classe"]), ("Libellé", r["libelle"]), ("Mois", r["mois"]), ("Mode", r["mode"]), ("Date", r["date"][:10])]:
        d.text((40, y), k, font=f(24), fill="#64748B"); d.text((250, y), str(v or "-"), font=f(26), fill="#0F172A"); y += 52
    d.rounded_rectangle((40, y + 10, 760, y + 100), 16, fill="#6366F1"); d.text((70, y + 38), fcfa(r["montant"]), font=f(40), fill="white")
    d.ellipse((560, 700, 740, 880), outline="#22C55E", width=8); d.text((600, 765), "PAYÉ", font=f(48), fill="#22C55E")
    q = qrcode.make(f"{BASE_URL}/verify/{r['recu']}").resize((200, 200)); im.paste(q, (40, 740))
    d.text((40, 960), "Fini de courir après les mensualités.", font=f(20), fill="#64748B")
    b = io.BytesIO(); im.save(b, "PNG"); b.seek(0); return b

@app.get("/recu/<code>.<fmt>")
def recu(code, fmt):
    r = recu_data(code)
    if not r: return "Reçu introuvable", 404
    if fmt == "pdf": return send_file(recu_pdf(r), mimetype="application/pdf", download_name=code + ".pdf")
    if fmt == "png": return send_file(recu_png(r), mimetype="image/png", download_name=code + ".png")
    k = r["montant"] // 1000
    return f"Reçu {k}k {r['nom']} - {r['classe']} - {r['matricule']} - {r['mode']} - {BASE_URL}/verify/{code}", 200, {"Content-Type": "text/plain; charset=utf-8"}

@app.get("/verify/<code>")
def verify(code):
    r = recu_data(code)
    ok = bool(r)
    box = (f"<p style='font-size:22px;color:#22C55E;font-weight:800'>✔ Reçu authentique</p><p><b>{r['nom']}</b> ({r['classe']})<br>{r['matricule']}</p>"
           f"<p style='font:700 24px monospace'>{fcfa(r['montant'])}</p><p>{r['libelle']} {r['mois'] or ''} · {r['mode']}<br>{r['date'].replace('T', ' ')}<br>{r['ecole']}</p>") if ok else \
          "<p style='font-size:22px;color:#EF4444;font-weight:800'>✖ Reçu introuvable ou faux</p>"
    return f"<!doctype html><meta charset=utf-8><meta name=viewport content='width=device-width,initial-scale=1'><body style='margin:0;background:#F8FAFC;font-family:Inter,sans-serif;display:grid;place-items:center;min-height:100vh'><div style='background:#fff;padding:32px;border-radius:20px;box-shadow:0 10px 40px #0001;text-align:center;max-width:380px'><img src='/static/logo.png' width=220><h3>Vérification {code}</h3>{box}</div>"

@app.get("/api/rapport.pdf")
@need()
def rapport():
    eid = session["ecole_id"]; d = db(); e = d.execute("SELECT nom FROM ecoles WHERE id=?", (eid,)).fetchone()
    rows = d.execute("SELECT p.date,p.recu,e.nom,e.classe,p.libelle,p.mode,p.montant FROM paiements p JOIN eleves e ON e.id=p.eleve_id WHERE p.ecole_id=? AND p.date LIKE ? ORDER BY p.date", (eid, datetime.now().strftime("%Y-%m") + "%")).fetchall()
    buf = io.BytesIO(); c = canvas.Canvas(buf, pagesize=A4); W, H = A4
    def head():
        c.setFillColorRGB(.059, .09, .165); c.rect(0, H - 80, W, 80, fill=1, stroke=0); c.drawImage(ImageReader(ICON()), 30, H - 72, 64, 64, mask="auto")
        c.setFillColorRGB(1, 1, 1); c.setFont("Helvetica-Bold", 16); c.drawString(110, H - 42, f"Rapport comptable – {e['nom']}"); c.setFont("Helvetica", 10); c.drawString(110, H - 60, datetime.now().strftime("Mois %m/%Y"))
        c.setFillColorRGB(0, 0, 0); return H - 105
    y = head(); tot = 0
    for r in rows:
        if y < 60: c.showPage(); y = head()
        c.setFont("Helvetica", 8); c.drawString(30, y, f"{r['date'][:16].replace('T',' ')}  {r['recu']}  {r['nom'][:24]}  {r['classe']}  {r['libelle'][:14]}  {r['mode']}"); c.setFont("Courier-Bold", 9); c.drawRightString(W - 30, y, fcfa(r["montant"])); tot += r["montant"]; y -= 15
    c.setFont("Courier-Bold", 13); c.drawRightString(W - 30, y - 15, "TOTAL : " + fcfa(tot)); c.showPage(); c.save(); buf.seek(0)
    return send_file(buf, mimetype="application/pdf", download_name="rapport.pdf")

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)), debug=False)
